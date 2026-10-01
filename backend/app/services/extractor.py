"""Extract clean, paragraph-structured text from PDF and DOCX files.

Output contract (the splitter relies on it):
- One Page per PDF page (number is 1-based) or a single Page(None, ...) for DOCX.
- Paragraphs inside Page.text are separated by "\n\n"; there are no other
  newlines. Text is NFKC-normalized, de-hyphenated, without bullet glyphs,
  page-number lines or running headers/footers.
- Page.continues is True when the page's last paragraph carries on at the top
  of the next page (mid-sentence or mid-word page break). The text keeps a
  trailing "-" if the break hyphenated a word; the splitter joins it.

Out of scope: OCR (text-less files raise NoTextError), multi-column layout,
tables. DOCX tables are skipped on purpose: read aloud cell by cell they are
noise, and python-docx keeps them out of Document.paragraphs anyway.
"""
import re
import statistics
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from docx import Document

X_TOLERANCE = 1.5  # lower if words glue together, raise if they split apart
HEADING_MAX_CHARS = 60
BULLETS = "•▪◦●■□‣⁃–—*·"

_PAGE_NUMBER_RE = re.compile(r"^\s*(page\s+)?\d{1,4}(\s*(/|of)\s*\d{1,4})?\s*$", re.IGNORECASE)
_BULLET_RE = re.compile(rf"^\s*[{re.escape(BULLETS)}]\s+")
# Hyphen at a line end between two word characters, e.g. "commer-\ncial".
_HYPHEN_BREAK_RE = re.compile(r"([\w-]*\w)-\n(\w[\w-]*)")
_DROP_CHARS_RE = re.compile("[­∗†‡§¶]")  # soft hyphen + footnote markers (never read aloud)
_TERMINAL_RE = re.compile(r"[.?!:][\"'”’)\]]*$")


class NoTextError(ValueError):
    """The document has no extractable text (e.g. a scanned PDF; OCR is out of scope)."""


@dataclass(frozen=True)
class Page:
    number: int | None  # 1-based PDF page number; None for DOCX
    text: str
    continues: bool = False  # last paragraph continues on the next page


def extract(path: str | Path) -> list[Page]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = _extract_pdf(path)
    elif suffix == ".docx":
        pages = _extract_docx(path)
    else:
        raise ValueError(f"Unsupported file type {suffix or '(none)'!r}: only .pdf and .docx are supported.")
    if not any(p.text for p in pages):
        raise NoTextError(f"No text found in {path.name}. Scanned documents (images only) are not supported.")
    return pages


# --------------------------------------------------------------------------- DOCX

def _extract_docx(path: Path) -> list[Page]:
    paragraphs = []
    for p in Document(path).paragraphs:
        text = _clean_inline(p.text)  # soft line breaks ("\n") and tabs become spaces
        if text:
            paragraphs.append(text)
    return [Page(None, "\n\n".join(paragraphs))]


# --------------------------------------------------------------------------- PDF

def _extract_pdf(path: Path) -> list[Page]:
    with pdfplumber.open(path) as pdf:
        # Drop rotated characters first (e.g. arXiv's vertical side stamp)
        raw = [(p.page_number, p.filter(_upright).extract_text_lines(x_tolerance=X_TOLERANCE, return_chars=False))
               for p in pdf.pages]
    repeated = _running_lines(raw)
    kept = [_kept_lines(lines, repeated) for _, lines in raw]
    pages = []
    for i, (number, _) in enumerate(raw):
        nxt = next((k for k in kept[i + 1:] if k), [])  # skip blank pages
        continues = bool(kept[i] and nxt) and _continues(kept[i], nxt[0])
        pages.append(Page(number, _paragraph_text(kept[i]), continues))
    return pages


def _upright(obj: dict) -> bool:
    return obj.get("object_type") != "char" or obj.get("upright", True)


def _line_key(text: str) -> str:
    """Header/footer identity: case-folded, digits masked ('Page 3' == 'Page 4')."""
    return re.sub(r"\d+", "#", _clean_inline(text).casefold())


def _running_lines(raw) -> set[str]:
    """Lines that repeat at the top or bottom of most pages (running headers/footers)."""
    if len(raw) < 3:
        return set()
    counts = Counter()
    for _, lines in raw:
        counts.update({_line_key(l["text"]) for l in lines[:2] + lines[-2:]})
    return {k for k, n in counts.items() if n >= max(3, len(raw) // 2 + 1) and k.strip("# ")}


def _kept_lines(lines: list[dict], repeated: set[str]) -> list[dict]:
    """Drop blank lines, page numbers and running headers/footers."""
    kept = []
    for i, l in enumerate(lines):
        text = _clean_inline(l["text"])
        if not text or _PAGE_NUMBER_RE.match(text):
            continue
        if (i < 2 or i >= len(lines) - 2) and _line_key(text) in repeated:
            continue
        kept.append({**l, "text": text})
    return kept


def _layout(kept: list[dict]) -> tuple[float, float]:
    """(right text edge, typical full line width) for the page."""
    return max(l["x1"] for l in kept), statistics.median(l["x1"] - l["x0"] for l in kept)


def _continues(kept: list[dict], next_first: dict) -> bool:
    """Does the page's last paragraph carry on at the top of the next page?"""
    last = kept[-1]
    if last["text"].endswith("-"):
        return True
    if _TERMINAL_RE.search(last["text"]):
        return False
    right_edge, full_width = _layout(kept)
    short = last["x1"] < right_edge - 0.15 * full_width
    return not (short and _is_heading(last, next_first))


def _paragraph_text(kept: list[dict]) -> str:
    if not kept:
        return ""
    gaps = [b["top"] - a["bottom"] for a, b in zip(kept, kept[1:]) if b["top"] > a["bottom"]]
    normal_gap = statistics.median(gaps) if gaps else 0
    right_edge, full_width = _layout(kept)

    paragraphs: list[list[str]] = []
    for i, l in enumerate(kept):
        text = l["text"]
        bullet = _BULLET_RE.match(text)
        if bullet:
            text = text[bullet.end():]
        prev = kept[i - 1] if i else None
        new_para = (
            prev is None
            or bullet
            or l["top"] - prev["bottom"] > max(normal_gap * 1.8, normal_gap + 3)  # vertical whitespace
            or _is_heading(prev, l)
            or _is_heading(l, kept[i + 1] if i + 1 < len(kept) else None)
            # previous line ended a sentence well before the right margin
            or (_TERMINAL_RE.search(prev["text"]) and prev["x1"] < right_edge - 0.15 * full_width)
        )
        if new_para:
            paragraphs.append([text])
        else:
            paragraphs[-1].append(text)
    return "\n\n".join(_join_lines(p) for p in paragraphs)


def _is_heading(line: dict | None, nxt: dict | None) -> bool:
    """Short line, no terminal punctuation, followed by a line starting with a capital."""
    if not line or not nxt:
        return False
    t = line["text"]
    return (
        len(t) < HEADING_MAX_CHARS
        and not t.endswith((".", "?", "!", ":", ",", ";", "-"))
        and nxt["text"][:1].isupper()
    )


def _join_lines(lines: list[str]) -> str:
    return dehyphenate("\n".join(lines)).replace("\n", " ")


def dehyphenate(text: str) -> str:
    """Join words hyphenated across line breaks ("commer-\\ncial" -> "commercial").

    If either side of the break already has a hyphen ("state-of-\\nthe-art",
    "English-\\nto-German") it is a compound, so the hyphen stays. Known
    trade-off: a two-part compound broken at its own hyphen ("mid-\\nrange")
    becomes "midrange".
    """
    return _HYPHEN_BREAK_RE.sub(lambda m: m[1] + ("-" if "-" in m[1] + m[2] else "") + m[2], text)


# --------------------------------------------------------------------------- shared

def _clean_inline(text: str) -> str:
    """NFKC (ligatures, full-width chars) + collapse all whitespace to single spaces."""
    text = unicodedata.normalize("NFKC", text)
    text = _DROP_CHARS_RE.sub("", text)
    return " ".join(text.split())
