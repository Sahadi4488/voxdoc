"""Generate the Day 3 test fixtures (committed; re-run only to change them).

    python tests\\fixtures\\make_fixtures.py      (from backend\\, needs reportlab)

sample.pdf  - 3 pages: running header, page-number footers, numbered headings
              without terminal punctuation, a bullet list, hyphenated line
              ends, literal ﬁ/ﬂ ligatures, a sentence crossing a page break,
              a >350-character sentence.
sample.docx - Heading 1/2, bullet list, abbreviations, a >350-character
              sentence, a soft line break, a small table.
scanned.pdf - one image-only page (no text layer) -> NoTextError.

Hyphenation points are marked with "|" in the source text below; the line
wrapper breaks a word there (adding "-") when it doesn't fit, like LaTeX does.
"""
import io
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).parent
FONTS = Path("C:/Windows/Fonts")
pdfmetrics.registerFont(TTFont("Arial", str(FONTS / "arial.ttf")))
pdfmetrics.registerFont(TTFont("Arial-Bold", str(FONTS / "arialbd.ttf")))

ABBREV = "Dr. Smith arrived at 3 p.m., e.g. this, The U.S. grew 3.5% in 2024."
LONG = (
    "Because a reader that waits for an entire chapter before speaking feels broken, "
    "VoxDoc synthesizes each sentence on its own, caches the result under a hash of the text, "
    "voice and speed, streams the first clip as soon as it is ready, and keeps generating the "
    "following sentences in the background, so that by the time the listener reaches the end "
    "of one clip the next one is usually already waiting on disk."
)

PDF_BLOCKS = [
    ("title", "Reading Documents Aloud with Small Neural Voices"),
    ("para", "Abstract. We describe VoxDoc, a doc|u|ment reader that turns PDF and Word ﬁles into "
             "natural speech on an ordinary laptop. The system runs a small text-to-speech model "
             "entirely on the CPU and highlights each word as it is spoken."),
    ("heading", "1. Introduction"),
    ("para", "Most people meet long documents at the worst possible moment: on a train, while "
             "cooking, or at the end of a tiring day. Lis|ten|ing is often easier than reading, yet "
             "com|mer|cial screen readers sound ro|bot|ic and cloud services send pri|vate ﬁles to "
             "someone else's server. A local reader avoids both prob|lems."),
    ("para", "Three goals shaped the design of the ﬁrst version:"),
    ("bullet", "Run on a mid-range laptop without a ded|i|cat|ed graphics card."),
    ("bullet", "Start speaking within a second of pressing play."),
    ("bullet", "Highlight the current word so the listener can follow along."),
    ("heading", "2. Method"),
    ("para", "The pipeline has four stages: ex|trac|tion, sen|tence split|ting, synthesis and play|back. "
             "Extraction pulls raw text out of each page and re|pairs the dam|age that typeset|ting "
             "leaves behind, such as words broken across lines and ligature characters like the "
             "one in efﬁcient work|ﬂow."),
    ("para", "Sentence splitting is harder than it looks. " + ABBREV + " A naive splitter would cut "
             "that line into ﬁve pieces. We rely on a sta|tis|ti|cal parser in|stead of reg|u|lar "
             "expressions, and we treat every paragraph as a hard boundary so that headings are "
             "never glued to the sentence that follows them."),
    ("heading", "2.1 Synthesis"),
    ("para", "Each sentence is synthesized with a state-of-the-art model of only eighty-two "
             "million pa|ram|e|ters. The model returns audio together with the time at which "
             "every word starts and ends, which is exactly what the in|ter|face needs for "
             "high|light|ing. Syn|the|sis runs about three times faster than real time on our test "
             "machine, so the listener never waits for the next sentence once playback has "
             "started, even when the doc|u|ment is long and the chosen voice is one of the "
             "slower ones."),
    ("para", LONG),
    ("heading", "3. Results"),
    ("para", "We tested the reader on forty doc|u|ments, including research papers, contracts and "
             "novels. Lis|ten|ers rated the voices as pleas|ant and the highlighting as accurate. "
             "The most common complaint con|cerned tables and ﬁgures, which the reader currently "
             "skips. Pages with running headers were handled correctly after we removed lines "
             "that repeat on every page."),
    ("heading", "3.1 Listening tests"),
    ("para", "Each par|tic|i|pant listened to the same three pas|sages with every pre|set voice and "
             "then chose a fa|vor|ite. Most people preferred the warm American voice for long "
             "reading ses|sions and the British voices for ﬁction. Several par|tic|i|pants asked "
             "for a slower default speed when the text was dense or tech|ni|cal, so the calm "
             "preset now starts at nine tenths of normal speed."),
    ("para", "We also measured how long a listener waits after pressing play. With an empty cache "
             "the ﬁrst sentence ar|rived in about one second; with a warm cache it ar|rived almost "
             "in|stant|ly, because the audio and its word timings are read straight from disk."),
    ("heading", "4. Conclusion"),
    ("para", "A small local model is good enough to read every|day doc|u|ments aloud. Future work "
             "will add support for scanned pages through optical char|ac|ter rec|og|ni|tion and "
             "will let listeners blend two voices into a new one."),
]

# Layout (points). Letter page, wide margins so the text spans 3 pages.
W, H = letter
LEFT, RIGHT, TOP, BOTTOM = 96, W - 96, H - 110, 130
STYLE = {  # font, size, leading, space before, space after
    "title": ("Arial-Bold", 16, 20, 0, 14),
    "heading": ("Arial-Bold", 13, 17, 10, 4),
    "para": ("Arial", 12, 16, 0, 10),
    "bullet": ("Arial", 12, 16, 0, 4),
}


def wrap(text, font, size, width):
    """Greedy wrap; breaks words at '|' hyphenation points (or real hyphens) when needed."""
    fits = lambda s: pdfmetrics.stringWidth(s, font, size) <= width  # noqa: E731
    lines, line = [], ""
    for word in text.split():
        while True:
            clean = word.replace("|", "")
            cand = f"{line} {clean}" if line else clean
            if fits(cand):
                line = cand
                break
            # Try the longest prefix ending at a '|' or '-' that still fits with a hyphen
            breaks = [i for i, c in enumerate(word) if c in "|-"]
            for i in reversed(breaks):
                head = word[:i].replace("|", "") + "-"
                if fits(f"{line} {head}" if line else head):
                    lines.append(f"{line} {head}" if line else head)
                    line, word = "", word[i + 1:]
                    break
            else:
                lines.append(line)
                line = ""
                continue
    if line:
        lines.append(line)
    return lines


def make_pdf(path):
    c = canvas.Canvas(str(path), pagesize=letter)
    page = 1

    def decorate():
        c.setFont("Arial", 9)
        c.drawString(LEFT, H - 50, "VoxDoc Fixture Paper - Draft for testing")
        c.drawCentredString(W / 2, 50, str(page))

    decorate()
    y = TOP
    for kind, text in PDF_BLOCKS:
        font, size, leading, before, after = STYLE[kind]
        x = LEFT + 22 if kind == "bullet" else LEFT
        y -= before
        for i, line in enumerate(wrap(text, font, size, RIGHT - x)):
            if y < BOTTOM:
                c.showPage()
                page += 1
                decorate()
                y = TOP
            c.setFont(font, size)
            if kind == "bullet" and i == 0:
                c.drawString(LEFT + 8, y, "•")
            c.drawString(x, y, line)
            y -= leading
        y -= after
    c.save()


def make_docx(path):
    d = Document()
    d.add_heading("VoxDoc Test Document", level=1)
    d.add_paragraph("This ﬁle checks how the extractor handles a Word document. It was generated "
                    "by a script so every edge case is present on purpose.")
    d.add_heading("Abbreviations and numbers", level=2)
    d.add_paragraph(ABBREV + " Prof. Lee costs approx. 5 dollars vs. 6 dollars at the store.")
    d.add_heading("A short list", level=2)
    for item in ["Extract the text.", "Split it into sentences.", "Read each sentence aloud."]:
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("A very long sentence", level=2)
    d.add_paragraph(LONG)
    p = d.add_paragraph("This paragraph has a soft line break")
    p.add_run().add_break()
    p.add_run("in the middle, which must become a space.")
    d.add_heading("A small table", level=2)
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([["Voice", "Accent", "RTF"], ["af_heart", "American", "0.33"],
                             ["bf_emma", "British", "0.31"]]):
        for col, val in enumerate(row):
            t.cell(r, col).text = val
    d.add_paragraph("The table above is skipped by the extractor. This is the last sentence.")
    d.save(path)


def make_scanned_pdf(path):
    img = Image.new("L", (1200, 300), 255)
    ImageDraw.Draw(img).text((40, 110), "This page is only an image of text.", fill=0,
                             font=ImageFont.truetype(str(FONTS / "arial.ttf"), 56))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawImage(ImageReader(io.BytesIO(buf.getvalue())), 72, H - 300, width=W - 144, height=(W - 144) / 4)
    c.save()


if __name__ == "__main__":
    make_pdf(HERE / "sample.pdf")
    make_docx(HERE / "sample.docx")
    make_scanned_pdf(HERE / "scanned.pdf")
    for f in ("sample.pdf", "sample.docx", "scanned.pdf"):
        print(f"{f:<12} {(HERE / f).stat().st_size / 1024:6.1f} KB")
